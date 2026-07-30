"""tableur — Adaptateur pour le gabarit xlsx d'ingestion T-R-C-O (§5.4), pour
une personne qui n'a pas à écrire de JSON à la main. Lit exactement les
onglets structurés produits par `scripts/generer_gabarit_ingestion.py`
("Tâches", "Ressources", "Précédences", "Compatibilités") — le 5e onglet du
gabarit ("Besoins additionnels") est du texte libre pour examen humain,
jamais traduit automatiquement.
"""

from __future__ import annotations

from collections.abc import Iterator
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from dsl.schema import (
    CompatibiliteRessourceTache,
    Contrainte,
    InstanceTRCO,
    MinimiserMakespan,
    Precedence,
    Ressource,
    Tache,
)

FEUILLES_REQUISES = ("Tâches", "Ressources", "Compatibilités")
FEUILLE_PRECEDENCES = "Précédences"


class ErreurFichierInvalide(Exception):
    """Fichier illisible ou onglet requis manquant — distinct d'une instance
    structurellement lue mais invalide au sens du DSL (ça, c'est une
    ValidationError Pydantic, même garde-fou que les autres canaux)."""


def _lignes(feuille: Worksheet, n_colonnes: int) -> Iterator[tuple[str, ...]]:
    """Lit les lignes de données (à partir de la ligne 2, en-tête exclu) —
    tolère la ligne d'exemple laissée telle quelle (elle sera alors traitée
    comme une vraie ligne, et un id dupliqué sera rejeté par InstanceTRCO,
    §6.7) et les lignes vides intercalées."""
    for ligne in feuille.iter_rows(min_row=2, max_col=n_colonnes, values_only=True):
        if not ligne or ligne[0] in (None, ""):
            continue
        yield tuple("" if v is None else str(v).strip() for v in ligne)


def _lire_taches(feuille: Worksheet) -> list[Tache]:
    return [Tache(id=id_, **({"nom": nom} if nom else {})) for id_, nom in _lignes(feuille, 2)]


def _lire_ressources(feuille: Worksheet) -> list[Ressource]:
    return [Ressource(id=id_, **({"nom": nom} if nom else {})) for id_, nom in _lignes(feuille, 2)]


def _lire_precedences(feuille: Worksheet) -> list[Contrainte]:
    return [Precedence(avant=avant, apres=apres) for avant, apres in _lignes(feuille, 2)]


def _lire_compatibilites(feuille: Worksheet) -> list[Contrainte]:
    compatibilites: list[Contrainte] = []
    for tache, ressource, duree in _lignes(feuille, 3):
        try:
            duree_jours = int(float(duree))
        except ValueError as erreur:
            raise ErreurFichierInvalide(
                f"onglet Compatibilités : durée invalide « {duree} » pour {tache}/{ressource} "
                "— doit être un nombre entier de jours."
            ) from erreur
        compatibilites.append(CompatibiliteRessourceTache(tache=tache, ressource=ressource, duree=duree_jours))
    return compatibilites


def traduire(contenu: bytes) -> InstanceTRCO:
    """Traduit un classeur xlsx (format du gabarit `docs/dsl/gabarit_ingestion_trco.xlsx`)
    en instance T-R-C-O. Lève `ErreurFichierInvalide` si le fichier n'est pas
    un xlsx lisible ou qu'un onglet requis manque ; `pydantic.ValidationError`
    si les données une fois lues ne forment pas une instance valide (id
    dupliqué, référence inconnue, tâche sans compatibilité...)."""
    try:
        classeur: Any = load_workbook(BytesIO(contenu), data_only=True, read_only=True)
    except Exception as erreur:
        raise ErreurFichierInvalide("fichier illisible : ce n'est pas un classeur Excel (.xlsx) valide") from erreur

    manquants = [f for f in FEUILLES_REQUISES if f not in classeur.sheetnames]
    if manquants:
        raise ErreurFichierInvalide(
            f"onglet(s) manquant(s) : {', '.join(manquants)} — utilisez le gabarit fourni "
            "(docs/dsl/gabarit_ingestion_trco.xlsx)."
        )

    taches = _lire_taches(classeur["Tâches"])
    ressources = _lire_ressources(classeur["Ressources"])
    precedences = _lire_precedences(classeur[FEUILLE_PRECEDENCES]) if FEUILLE_PRECEDENCES in classeur.sheetnames else []
    compatibilites = _lire_compatibilites(classeur["Compatibilités"])

    return InstanceTRCO(
        taches=taches,
        ressources=ressources,
        contraintes=[*precedences, *compatibilites],
        objectifs=[MinimiserMakespan()],
    )
