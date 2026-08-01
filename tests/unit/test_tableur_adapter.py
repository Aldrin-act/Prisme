"""Couche 1 (§6.1) : `adapters/tableur/traducteur.py` — la version xlsx
(un seul classeur à onglets) du gabarit d'ingestion T-R-C-O, miroir de
`adapters/csv_import/` (trois fichiers séparés). Construit les classeurs de
test en mémoire avec `openpyxl` plutôt que de dépendre du gabarit réel
(`docs/dsl/gabarit_ingestion_trco.xlsx`), pour rester indépendant de son
contenu d'exemple.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from adapters.tableur import ErreurFichierInvalide, traduire


def _classeur(
    taches: list[tuple],
    ressources: list[tuple],
    compatibilites: list[tuple],
    precedences: list[tuple] | None = None,
    competences_requises: list[tuple] | None = None,
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Tâches")
    ws.append(["id *", "nom", "duree_estimee_jours"])
    for ligne in taches:
        ws.append(list(ligne))

    ws = wb.create_sheet("Ressources")
    ws.append(["id *", "nom", "competences"])
    for ligne in ressources:
        ws.append(list(ligne))

    if precedences is not None:
        ws = wb.create_sheet("Précédences")
        ws.append(["tache_avant *", "tache_apres *"])
        for ligne in precedences:
            ws.append(list(ligne))

    if competences_requises is not None:
        ws = wb.create_sheet("Compétences requises")
        ws.append(["tache *", "competence *"])
        for ligne in competences_requises:
            ws.append(list(ligne))

    ws = wb.create_sheet("Compatibilités")
    ws.append(["tache *", "ressource *", "duree_jours *"])
    for ligne in compatibilites:
        ws.append(list(ligne))

    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


TACHES = [("T1", "Decoupe", ""), ("T2", "Assemblage", "")]
RESSOURCES = [("R1", "Decoupeuse", "")]
COMPATIBILITES = [("T1", "R1", 10), ("T2", "R1", 15)]


def test_traduire_lit_les_onglets_de_base() -> None:
    contenu = _classeur(TACHES, RESSOURCES, COMPATIBILITES, precedences=[("T1", "T2")])

    instance = traduire(contenu)

    assert [t.id for t in instance.taches] == ["T1", "T2"]
    assert instance.taches[0].nom == "Decoupe"
    assert [r.id for r in instance.ressources] == ["R1"]
    assert len(instance.contraintes) == 3
    assert instance.objectifs[0].type == "minimiser_makespan"


def test_onglet_precedences_absent_tolere() -> None:
    """Comme en CSV, les précédences restent optionnelles."""
    contenu = _classeur(TACHES, RESSOURCES, COMPATIBILITES, precedences=None)
    instance = traduire(contenu)
    assert not any(c.type == "precedence" for c in instance.contraintes)


def test_onglet_competences_requises_absent_tolere() -> None:
    """Un classeur sans onglet "Compétences requises" (gabarit plus ancien,
    ou simplement inutilisé) doit se comporter exactement comme avant son
    introduction — aucune erreur, aucune dérivation."""
    contenu = _classeur(TACHES, RESSOURCES, COMPATIBILITES, competences_requises=None)
    instance = traduire(contenu)
    assert not any(c.type == "competence_requise" for c in instance.contraintes)


def test_onglet_requis_manquant_leve_erreur_fichier_invalide() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Tâches")
    ws.append(["id *", "nom", "duree_estimee_jours"])
    ws.append(["T1", "Decoupe", ""])
    tampon = BytesIO()
    wb.save(tampon)

    with pytest.raises(ErreurFichierInvalide, match="onglet"):
        traduire(tampon.getvalue())


def test_fichier_illisible_leve_erreur_fichier_invalide() -> None:
    with pytest.raises(ErreurFichierInvalide, match="classeur Excel"):
        traduire(b"ceci n'est pas un fichier xlsx")


def test_duree_non_numerique_leve_erreur_fichier_invalide() -> None:
    contenu = _classeur(TACHES, RESSOURCES, [("T1", "R1", "pas-un-nombre")])
    with pytest.raises(ErreurFichierInvalide, match="durée invalide"):
        traduire(contenu)


def test_tache_sans_compatibilite_rejetee_par_le_garde_fou_dsl() -> None:
    """Le garde-fou (§6.7) tranche pareil, quel que soit le canal d'ingestion :
    T2 n'a aucune compatibilité ici."""
    contenu = _classeur(TACHES, RESSOURCES, [("T1", "R1", 10)])
    with pytest.raises(ValidationError):
        traduire(contenu)


# --- Compatibilité dérivée par compétence (onglet "Compétences requises") ---


def test_compatibilite_derivee_pour_chaque_ressource_competente() -> None:
    taches = [("T1", "Decoupe", 25)]
    ressources = [("R1", "Decoupeuse", "decoupe;affutage"), ("R2", "Assembleuse", "assemblage")]

    instance = traduire(_classeur(taches, ressources, compatibilites=[], competences_requises=[("T1", "decoupe")]))

    compatibilites = [c for c in instance.contraintes if c.type == "compatibilite_ressource_tache"]
    assert len(compatibilites) == 1
    assert compatibilites[0].tache == "T1"
    assert compatibilites[0].ressource == "R1"
    assert compatibilites[0].duree == 25
    # R2 n'a pas la compétence requise — jamais rendue compatible.
    assert not any(c.ressource == "R2" for c in compatibilites)
    # La compétence requise elle-même reste dans l'instance (garde-fou DSL, §6.7).
    assert any(c.type == "competence_requise" and c.tache == "T1" for c in instance.contraintes)


def test_compatibilite_derivee_exige_toutes_les_competences_requises() -> None:
    """Une tâche à plusieurs compétences requises n'est compatible qu'avec une
    ressource qui les couvre TOUTES — même exigence que le garde-fou DSL,
    pas une simple intersection non vide."""
    taches = [("T1", "", 25)]
    ressources = [("R1", "", "decoupe"), ("R2", "", "decoupe;affutage")]
    competences_requises = [("T1", "decoupe"), ("T1", "affutage")]

    instance = traduire(
        _classeur(taches, ressources, compatibilites=[], competences_requises=competences_requises)
    )

    compatibilites = [c for c in instance.contraintes if c.type == "compatibilite_ressource_tache"]
    assert [c.ressource for c in compatibilites] == ["R2"]


def test_compatibilite_derivee_sans_duree_estimee_leve_erreur_fichier_invalide() -> None:
    taches = [("T1", "", "")]
    ressources = [("R1", "", "decoupe")]

    with pytest.raises(ErreurFichierInvalide, match="durée estimée manquante"):
        traduire(_classeur(taches, ressources, compatibilites=[], competences_requises=[("T1", "decoupe")]))


def test_compatibilite_explicite_et_derivee_par_competence_coexistent() -> None:
    """Une même tâche peut combiner compatibilité déclarée à la main (onglet
    Compatibilités) et compatibilité dérivée par compétence — les deux
    mécanismes s'additionnent, même règle qu'en CSV."""
    taches = [("T1", "", 25)]
    ressources = [("R1", "", "decoupe"), ("R2", "", "decoupe")]
    compatibilites = [("T1", "R2", 40)]
    competences_requises = [("T1", "decoupe")]

    instance = traduire(_classeur(taches, ressources, compatibilites, competences_requises=competences_requises))

    obtenues = {
        (c.tache, c.ressource, c.duree) for c in instance.contraintes if c.type == "compatibilite_ressource_tache"
    }
    assert obtenues == {("T1", "R2", 40), ("T1", "R1", 25)}
